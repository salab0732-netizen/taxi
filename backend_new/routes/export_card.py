"""
تصدير Excel — «البطاقة الجماعية لسائقي سيارات الأجرة»
─────────────────────────────────────────────────────
• سطر واحد لكل سائق = نسخة جماعية من «بطاقة معلومات سائق سيارة الأجرة»
• الوضعية الحالية فقط (السجلات is_current=1 / الباب النشط) — عند أي تغيير
  يُؤرشف القديم ويصبح الجديد هو الحالي، فيظهر تلقائياً في نفس السطر.
• التاريخ الكامل لكل سائق مكانه الشهادة التاريخية الفردية، لا هذا الملف.
"""
import json
from datetime import datetime, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NA = "غير مسجّل"
ACT = {"فردية_حضرية": "فردية حضرية", "جماعية_حضرية": "جماعية حضرية",
       "مابين_البلديات": "ما بين البلديات", "مابين_الولايات": "ما بين الولايات"}

# (عنوان القسم، لون، [(عنوان العمود، مفتاح، عرض)])
SECTIONS = [
    ("السائق", "125950", [
        ("#", "_n", 5), ("الحالة", "statut", 12), ("نمط النشاط", "activity", 16),
    ]),
    ("الهوية", "1D4ED8", [
        ("اللقب", "nom_ar", 14), ("الاسم", "prenom_ar", 14), ("Nom", "nom_fr", 14), ("Prénom", "prenom_fr", 14),
        ("تاريخ الميلاد", "date_naissance", 12), ("مكان الميلاد", "lieu_naissance_ar", 16),
        ("NIN", "nin", 21), ("العنوان", "adresse", 26), ("الهاتف", "telephone", 13),
    ]),
    ("رخصة السياقة", "0891B2", [
        ("رقم الرخصة", "lic_num", 13), ("الفئات", "lic_cats", 9), ("تاريخ الإصدار", "lic_from", 12),
        ("تاريخ الانتهاء", "lic_to", 12), ("جهة الإصدار", "lic_where", 18),
    ]),
    ("المركبة", "7C3AED", [
        ("رقم التسجيل", "veh_num", 15), ("الصنف", "veh_marque", 12), ("الطراز", "veh_type", 13),
        ("الرقم التسلسلي", "veh_serie", 20), ("الوقود", "veh_energie", 9), ("المقاعد", "veh_places", 8),
        ("سنة أول استعمال", "veh_year", 10),
    ]),
    ("الباب والمستفيد", "D97706", [
        ("رقم الباب", "door", 9), ("صفة الاستغلال", "door_mode", 11), ("بلدية الإلحاق", "door_commune", 13), ("نوع القرار", "dec_type", 15),
        ("رقم القرار", "dec_num", 10), ("تاريخ القرار", "dec_date", 12),
        ("المستفيد", "ben_name", 20), ("NIN المستفيد", "ben_nin", 21), ("صفة المستفيد", "ben_sifa", 14),
    ]),
    ("عقد الكراء", "059669", [
        ("رقم العقد", "rc_num", 17), ("تاريخ التحرير", "rc_from", 12), ("تاريخ الانتهاء", "rc_to", 12),
        ("الإيجار الشهري (دج)", "rc_rent", 12),
    ]),
    ("المناوب", "6366F1", [
        ("المناوب", "dep_name", 20), ("NIN المناوب", "dep_nin", 21), ("هاتف المناوب", "dep_tel", 13),
        ("رخصة المناوب", "dep_permis", 13), ("فئاتها", "dep_cats", 9), ("انتهاء رخصته", "dep_to", 12),
        ("عقد المناوب", "dc_num", 17), ("انتهاء العقد", "dc_to", 12),
        ("رخصة السائق الإضافي", "prm_num", 17), ("صلاحيتها", "prm_to", 12),
    ]),
]
DATE_WARN_KEYS = ("lic_to", "rc_to", "dep_to", "dc_to", "prm_to")


def _txt(v):
    if v is None or str(v).strip() in ("", "None"):
        return None
    s = str(v).strip()
    if s.startswith("["):
        try:
            a = json.loads(s)
            if isinstance(a, list):
                return "، ".join(map(str, a)) or None
        except Exception:
            pass
    if "_" in s and any("؀" <= ch <= "ۿ" for ch in s):
        s = s.replace("_", " ")
    return s


def _d10(v):
    return str(v)[:10] if v else None


def _collect(conn):
    rows = []
    drivers = conn.execute("""
        SELECT d.* FROM drivers d JOIN accounts a ON a.id=d.account_id AND a.role='driver'
        ORDER BY d.id
    """).fetchall()
    # تحميل مسبق لكل جدول مرة واحدة (سريع حتى مع آلاف السائقين)
    def by_driver(sql, key="driver_id"):
        m = {}
        for r in conn.execute(sql).fetchall():
            m[r[key]] = r          # ORDER BY id ⇒ الأحدث يبقى
        return m
    LIC = by_driver("SELECT * FROM driver_licenses WHERE is_current=1 ORDER BY id")
    VEH = by_driver("SELECT * FROM vehicles WHERE is_current=1 ORDER BY id")
    DOOR = by_driver("""SELECT dl.*, b.nom_ar bn, b.prenom_ar bp, b.nin bnin, b.sifa FROM door_licenses dl
                        LEFT JOIN beneficiaries b ON b.id=dl.beneficiary_id
                        WHERE dl.is_active=1 AND dl.current_driver_id IS NOT NULL ORDER BY dl.id""", "current_driver_id")
    RC = by_driver("SELECT * FROM rental_contracts WHERE is_current=1 ORDER BY id")
    DEP = by_driver("SELECT * FROM deputies WHERE is_current=1 ORDER BY id")
    DC = by_driver("SELECT * FROM deputy_contracts WHERE is_current=1 ORDER BY id")
    ACTS = by_driver("SELECT * FROM activity WHERE is_current=1 ORDER BY id")
    PRM = by_driver("SELECT * FROM deputy_permits WHERE is_current=1 ORDER BY id")
    for d in drivers:
        d = dict(d); i = d["id"]
        lic, veh, door, rc = LIC.get(i), VEH.get(i), DOOR.get(i), RC.get(i)
        dep, dc, act, prm = DEP.get(i), DC.get(i), ACTS.get(i), PRM.get(i)
        g = lambda r, k: (dict(r).get(k) if r else None)
        rows.append({
            "statut": _txt(d.get("statut")) or "نشط",
            "activity": ACT.get(g(act, "activity_type"), _txt(g(act, "activity_type"))),
            "nom_ar": d.get("nom_ar"), "prenom_ar": d.get("prenom_ar"),
            "nom_fr": d.get("nom_fr"), "prenom_fr": d.get("prenom_fr"),
            "date_naissance": d.get("date_naissance"), "lieu_naissance_ar": d.get("lieu_naissance_ar"),
            "nin": d.get("nin"), "adresse": d.get("adresse"), "telephone": d.get("telephone"),
            "lic_num": g(lic, "num_permis"), "lic_cats": _txt(g(lic, "categories")),
            "lic_from": _d10(g(lic, "date_delivrance")), "lic_to": _d10(g(lic, "date_expiration")),
            "lic_where": g(lic, "lieu_delivrance") or g(lic, "wilaya_delivrance"),
            "veh_num": g(veh, "num_immatriculation"), "veh_marque": g(veh, "marque"),
            "veh_type": g(veh, "type_vehicule"), "veh_serie": g(veh, "num_serie"),
            "veh_energie": g(veh, "energie"), "veh_places": g(veh, "nb_places"),
            "veh_year": g(veh, "annee_circulation"),
            "door": g(door, "door_number"),
            "door_mode": (g(door, "exploitation_mode") or "مستأجر") if door else None,
            "door_commune": g(door, "exploitation_commune") or g(door, "wilaya"),
            "dec_type": _txt(g(door, "decision_type")), "dec_num": g(door, "decision_number"),
            "dec_date": _d10(g(door, "decision_date")),
            "ben_name": (f"{g(door,'bp') or ''} {g(door,'bn') or ''}".strip() or None) if door else None,
            "ben_nin": g(door, "bnin"), "ben_sifa": _txt(g(door, "sifa")),
            **({"rc_num": "مستفيد", "rc_from": "مستفيد", "rc_to": "مستفيد", "rc_rent": "مستفيد"}
               if door and g(door, "exploitation_mode") == "مستفيد" else
               {"rc_num": g(rc, "contract_number"), "rc_from": _d10(g(rc, "contract_date")),
                "rc_to": _d10(g(rc, "end_date")), "rc_rent": g(rc, "monthly_rent") or None}),
            "dep_name": (f"{g(dep,'prenom_ar') or ''} {g(dep,'nom_ar') or ''}".strip() or None) if dep else None,
            "dep_nin": g(dep, "nin"), "dep_tel": g(dep, "telephone"), "dep_permis": g(dep, "num_permis"),
            "dep_cats": _txt(g(dep, "categories_permis")), "dep_to": _d10(g(dep, "date_expiration_permis")),
            "dc_num": g(dc, "contract_number"), "dc_to": _d10(g(dc, "end_date")),
            "prm_num": g(prm, "permit_number"), "prm_to": _d10(g(prm, "expiry_date")),
        })
    return rows


def build_workbook(conn):
    rows = _collect(conn)
    today = datetime.now().date()
    soon = today + timedelta(days=30)
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = _canevas_sheet(conn)   # ورقة 1: الملخص الرسمي (Canevas n°03 Activité Taxis)

    # ════════ ورقة 2: البطاقة الجماعية ════════
    ws = wb.create_sheet("البطاقة الجماعية"); ws.sheet_view.rightToLeft = True
    col = 1
    cols = []
    for title, color, fields in SECTIONS:
        start = col
        for (hdr, key, width) in fields:
            h = ws.cell(2, col, hdr)
            h.font = Font(bold=True, color="FFFFFF", size=10)
            h.fill = PatternFill("solid", fgColor=color); h.alignment = center; h.border = border
            ws.column_dimensions[get_column_letter(col)].width = width
            cols.append((key, color))
            col += 1
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)
        t = ws.cell(1, start, title)
        t.font = Font(bold=True, color="FFFFFF", size=11)
        t.fill = PatternFill("solid", fgColor=color); t.alignment = center
    ws.row_dimensions[2].height = 32

    stop_fill = PatternFill("solid", fgColor="FEE2E2")
    temp_fill = PatternFill("solid", fgColor="FEF3C7")
    zebra = PatternFill("solid", fgColor="F8FAFC")
    na_font = Font(color="9CA3AF", italic=True, size=9)
    for n, r in enumerate(rows, 1):
        ri = n + 2
        fill = stop_fill if r["statut"] == "توقف نهائي" else temp_fill if r["statut"] == "توقف مؤقت" else (zebra if n % 2 == 0 else None)
        for ci, (key, _) in enumerate(cols, 1):
            val = n if key == "_n" else r.get(key)
            cell = ws.cell(ri, ci, val if val not in (None, "") else NA)
            cell.border = border; cell.alignment = Alignment(horizontal="center", vertical="center")
            if fill: cell.fill = fill
            if val in (None, ""):
                cell.font = na_font
            elif key in DATE_WARN_KEYS and isinstance(val, str) and val[:1].isdigit():
                if val < today.isoformat():
                    cell.font = Font(bold=True, color="DC2626")
                elif val <= soon.isoformat():
                    cell.font = Font(bold=True, color="D97706")
            elif key == "statut" and val != "نشط":
                cell.font = Font(bold=True, color="B91C1C")
    ws.freeze_panes = "E3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{max(2, len(rows) + 2)}"

    lg = len(rows) + 4
    ws.cell(lg, 1, "دليل الألوان:").font = Font(bold=True)
    ws.cell(lg, 2, "توقف نهائي").fill = stop_fill
    ws.cell(lg, 3, "توقف مؤقت").fill = temp_fill
    c = ws.cell(lg, 4, "تاريخ منتهٍ"); c.font = Font(bold=True, color="DC2626")
    c = ws.cell(lg, 5, "ينتهي خلال 30 يوماً"); c.font = Font(bold=True, color="D97706")
    ws.cell(lg + 1, 1, "هذا الملف يعرض الوضعية الحالية فقط — التاريخ الكامل لكل سائق في «الشهادة التاريخية».").font = Font(italic=True, color="6B7280")

    wb.active = 0
    return wb


# ════════════════════════════════════════════════════════════════
# الملخص الرسمي — نموذج «Canevas n°03 Activité Taxis» (مديرية النقل)
# القالب: backend_new/templates/canevas_taxi_template.xlsx
# تُملأ الخانات من الوضعية الحالية في قاعدة البيانات؛ الصيغ تبقى حيّة.
# ════════════════════════════════════════════════════════════════
import os
from pathlib import Path

TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "canevas_taxi_template.xlsx"
MOUDJ = ("مجاهد", "ابن_مجاهد", "ابن_شهيد", "أرملة_مجاهد", "ابن مجاهد", "ابن شهيد", "أرملة مجاهد")
# صف كل نمط استغلال في القالب
MODE_ROW = {"فردية_حضرية": 20, "جماعية_حضرية": 21, "مابين_البلديات": 23, "مابين_الولايات": 24}
AGE_COLS = "BCDEFGHI"   # [0,5[ [5,10[ ... [35,+[


def vehicle_year(annee, plate):
    """سنة أول استعمال: من البطاقة الرمادية، وإلا من رقم التسجيل الجزائري.
    الصيغة NNNNN-CYY-WW : C = صنف المركبة، YY = آخر رقمين من سنة أول تسجيل، WW = رمز الولاية
    مثال: 00998-116-32 → 1 (سياحية) + 16 → 2016، ولاية 32 (البيض)"""
    import re
    m = re.match(r"\s*(\d{4})", str(annee or ""))
    if m:
        return int(m.group(1))
    m = re.search(r"\d+\s*[-/ ]\s*(\d)(\d{2})\s*[-/ ]\s*\d{2}", str(plate or ""))
    if not m:
        return None
    yy = int(m.group(2))
    now = datetime.now().year
    y = 2000 + yy
    return y if y <= now else 1900 + yy


def _canevas_sheet(conn):
    try:
        from routes.company import expire_contracts
        expire_contracts(conn)
    except Exception:
        pass
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.worksheets[0]
    now = datetime.now()

    # ── الرأس ──
    ws["A2"] = "WILAYA DE : EL BAYADH"
    ws["A3"] = f"ANNEE : {now.year}"
    ws["A4"] = f"TRIMESTRE : {(now.month - 1) // 3 + 1:02d}"

    # ── 1) الرخص المسندة والمستغلة ──
    doors = conn.execute("""
        SELECT dl.door_number, dl.is_active, dl.current_driver_id, dl.exploitation_mode, b.sifa
        FROM door_licenses dl
        LEFT JOIN beneficiaries b ON b.id = dl.beneficiary_id
        WHERE dl.id IN (SELECT MAX(id) FROM door_licenses GROUP BY door_number)
    """).fetchall()
    is_partic = lambda x: (x["sifa"] or "") == "خاص"
    # Licences attribuées: ذوو الحقوق / خواص
    attrib_p = sum(1 for x in doors if is_partic(x))
    attrib_m = len(doors) - attrib_p
    # Licences exploitées (أبواب مرتبطة حالياً بسائق):
    #   particuliers = رخصة خاص (يستغلها بنفسه أو يكريها) ؛
    #   Bénéficiaires directs = ذو حق يستغل رخصته بنفسه (صفة الاستغلال «مستفيد») ؛
    #   Locataires = ذو حق أكرى رخصته لسائق مستأجر
    exploited = [x for x in doors if x["is_active"] and x["current_driver_id"]]
    partic = sum(1 for x in exploited if is_partic(x))
    direct = sum(1 for x in exploited if not is_partic(x) and (x["exploitation_mode"] or "مستأجر") == "مستفيد")
    locat = len(exploited) - partic - direct
    ws["B10"], ws["C10"] = attrib_m, attrib_p
    ws["E10"], ws["F10"], ws["G10"] = direct, partic, locat
    ws["I10"] = "=IFERROR(H10*100/D10,0)"

    # ── 2) عدد السائقين: المستغلون (نشط + باب نشط) والمناوبون (عقد مناوب ساري) ──
    active_ids = [r["id"] for r in conn.execute("""
        SELECT DISTINCT d.id FROM drivers d
        JOIN accounts a ON a.id = d.account_id AND a.role = 'driver'
        JOIN door_licenses dl ON dl.current_driver_id = d.id AND dl.is_active = 1
        WHERE d.statut = 'نشط'
    """).fetchall()]
    ws["A15"] = len(active_ids)
    ws["B15"] = conn.execute(
        f"SELECT COUNT(*) FROM deputy_contracts WHERE is_current=1 AND driver_id IN ({','.join('?'*len(active_ids)) or 'NULL'})",
        active_ids).fetchone()[0] if active_ids else 0

    # ── 3) توزيع الحظيرة حسب العمر ونمط الاستغلال ──
    for r in range(20, 26):
        for c in AGE_COLS:
            ws[f"{c}{r}"] = 0
    if active_ids:
        q = f"""
            SELECT v.annee_circulation, v.num_immatriculation, a.activity_type FROM vehicles v
            LEFT JOIN activity a ON a.driver_id = v.driver_id AND a.is_current = 1
            WHERE v.is_current = 1 AND v.driver_id IN ({','.join('?'*len(active_ids))})
        """
        for row in conn.execute(q, active_ids).fetchall():
            y = vehicle_year(row["annee_circulation"], row["num_immatriculation"])
            if not y:
                continue
            age = now.year - y
            r = MODE_ROW.get(row["activity_type"] or "فردية_حضرية", 20)
            col = AGE_COLS[min(max(age, 0) // 5, 7)]
            ws[f"{col}{r}"] = (ws[f"{col}{r}"].value or 0) + 1
    # تصحيح القالب: مجموع كل عمود يشمل الصفوف 20→25 (كان يستثني سيارات الأجرة الفردية)
    for c in AGE_COLS:
        ws[f"{c}26"] = f"=SUM({c}20:{c}25)"
    for r in range(20, 27):
        f = ws[f"K{r}"].value
        if isinstance(f, str) and f.startswith("=") and "IFERROR" not in f:
            ws[f"K{r}"] = "=IFERROR(" + f[1:] + ",0)"

    # ── 4) و 5) شركات سيارات الأجرة — من فضاء الشركات ──
    #   الحظيرة = مركبات الشركة حسب سنة أول استعمال
    #   السائقون = الحائزون على رخصة سائق أجير سارية
    #   Opérationnelle (Date) = تاريخ اعتماد الشركة
    companies = conn.execute("""
        SELECT c.id, COALESCE(NULLIF(c.nom_fr,''), c.nom_ar) AS nom, c.date_agrement
        FROM companies c JOIN accounts a ON a.id = c.account_id
        WHERE a.role = 'company' AND a.is_active = 1 AND c.nom_ar IS NOT NULL
        ORDER BY c.id
    """).fetchall()
    rows = []
    for co in companies:
        ages = [0] * 8
        for v in conn.execute("SELECT annee_circulation, num_immatriculation FROM company_vehicles WHERE company_id=?", (co["id"],)):
            y = vehicle_year(v["annee_circulation"], v["num_immatriculation"])
            if not y:
                continue
            age = now.year - y
            ages[min(max(age, 0) // 5, 7)] += 1
        drivers = conn.execute("""SELECT COUNT(DISTINCT driver_id) FROM company_hire_requests
                                  WHERE company_id=? AND statut='مقبول' AND is_current=1""", (co["id"],)).fetchone()[0]
        rows.append({"nom": co["nom"], "ages": ages, "drivers": drivers, "date": co["date_agrement"]})
    # 9 أسطر في القالب (37→45): إن زادت الشركات تُجمع البقية في السطر الأخير
    if len(rows) > 9:
        rest = rows[8:]
        rows = rows[:8] + [{"nom": f"Autres ({len(rest)} sociétés)",
                            "ages": [sum(x["ages"][k] for x in rest) for k in range(8)],
                            "drivers": sum(x["drivers"] for x in rest), "date": None}]
    for idx, r in enumerate(range(37, 46)):
        f = ws[f"L{r}"].value
        if isinstance(f, str) and f.startswith("=") and "IFERROR" not in f:
            ws[f"L{r}"] = "=IFERROR(" + f[1:] + ",0)"
        if idx < len(rows):
            x = rows[idx]
            ws[f"A{r}"] = idx + 1
            ws[f"B{r}"] = x["nom"]
            for k, c in enumerate("CDEFGHIJ"):
                ws[f"{c}{r}"] = x["ages"][k]
            ws[f"M{r}"] = x["drivers"]
            try:
                ws[f"N{r}"] = datetime.strptime(x["date"][:10], "%Y-%m-%d") if x["date"] else None
                if x["date"]:
                    ws[f"N{r}"].number_format = "DD/MM/YYYY"
            except (TypeError, ValueError):
                ws[f"N{r}"] = x["date"]
        else:
            ws[f"A{r}"] = None
            ws[f"B{r}"] = None
            for c in "CDEFGHIJ":
                ws[f"{c}{r}"] = None
            ws[f"M{r}"] = None
            ws[f"N{r}"] = None
    ws["B30"] = len(companies)
    ws["B31"] = "=M46"
    ws["B46"] = "=COUNTA(B37:B45)"
    for cell in ("L46", "K51", "K52", "K53"):
        f = ws[cell].value
        if isinstance(f, str) and f.startswith("=") and "IFERROR" not in f:
            ws[cell] = "=IFERROR(" + f[1:] + ",0)"
    for c in AGE_COLS:
        ws[f"{c}54"] = f"=IFERROR({c}53*100/J53,0)"

    ws["A56"] = (f"Généré automatiquement le {now.strftime('%d/%m/%Y %H:%M')} — sections 1 à 5 calculées depuis la base "
                 "(parc selon l'année de 1ère mise en circulation, sinon déduite du n° d'immatriculation ; conducteurs = titulaires d'une licence de chauffeur salarié valide).")
    ws["A56"].font = Font(italic=True, size=9, color="6B7280")
    ws.title = "Canevas 03 - Taxis"
    return wb
