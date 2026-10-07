"""
routes/company.py — فضاء شركات سيارات الأجرة
هوية الشركة + المركبات + السائقون وربطهم بالمركبات
"""
from flask import Blueprint, request, jsonify, Response
from datetime import datetime
from html import escape
from functools import wraps
from database import get_db
from utils import require_auth, require_admin, save_image
from database import generate_number

company_bp = Blueprint("company", __name__)

PROFILE_FIELDS = [
    "nom_ar", "nom_fr", "adresse", "commune", "wilaya", "telephone", "telephone2",
    "email", "registre_commerce", "rc_date", "num_fiscal", "num_agrement", "date_agrement",
    "gerant_nom_ar", "gerant_prenom_ar", "gerant_nom_fr", "gerant_prenom_fr",
    "gerant_date_naissance", "gerant_lieu_naissance", "gerant_nin", "gerant_adresse",
    "representant_nom", "representant_fonction", "notes",
]

VEHICLE_FIELDS = [
    "num_immatriculation", "num_precedent", "marque", "type_vehicule", "num_serie",
    "genre", "carrosserie", "energie", "puissance", "nb_places", "poids_total",
    "charge_utile", "annee_circulation", "date_delivrance", "lieu_delivrance",
    "wilaya_delivrance", "quittance_num", "quittance_montant", "quittance_date",
    "proprietaire_nom_ar", "proprietaire_prenom_ar", "proprietaire_nom",
    "proprietaire_prenom", "proprietaire_adresse", "proprietaire_wilaya",
]

DRIVER_FIELDS = [
    "nom_ar", "prenom_ar", "nom_fr", "prenom_fr", "date_naissance", "lieu_naissance",
    "nin", "telephone", "num_permis", "date_delivrance", "date_expiration",
    "lieu_delivrance", "wilaya_delivrance", "categories",
]


def expire_contracts(conn):
    """العقود التي تجاوزت تاريخ نهايتها تنتهي تلقائياً (يُستدعى قبل أي قراءة/تحقق)"""
    today = datetime.now().strftime("%Y-%m-%d")
    conn.execute("""UPDATE company_hire_requests
                    SET is_current=0, end_reason='انتهاء_المدة', terminated_at=?,
                        statut = CASE WHEN statut='جديد' THEN 'ملغى' ELSE statut END
                    WHERE is_current=1 AND contract_end IS NOT NULL AND contract_end < ?""",
                 (today + " 00:00:00", today))
    conn.commit()


def _active_contract(conn, driver_id):
    """عقد توظيف ساري (قيد الدراسة أو مقبول) لهذا السائق"""
    return conn.execute("""SELECT * FROM company_hire_requests
                           WHERE driver_id=? AND is_current=1 AND statut IN ('جديد','مقبول')""",
                        (driver_id,)).fetchone()


def _clean(v):
    if isinstance(v, list):
        v = ",".join(str(x) for x in v)
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def require_company(f):
    """يتحقق أن الحساب حساب شركة ويمرّر company_id (ينشئ سجل الشركة عند الحاجة)"""
    @require_auth
    @wraps(f)
    def decorated(*args, account=None, **kwargs):
        if account["role"] != "company":
            return jsonify({"error": "غير مخوّل"}), 403
        with get_db() as conn:
            expire_contracts(conn)
            co = conn.execute("SELECT id FROM companies WHERE account_id=?",
                              (account["id"],)).fetchone()
            if not co:
                conn.execute("INSERT INTO companies (account_id) VALUES (?)", (account["id"],))
                conn.commit()
                co = conn.execute("SELECT id FROM companies WHERE account_id=?",
                                  (account["id"],)).fetchone()
        return f(*args, account=account, company_id=co["id"], **kwargs)
    return decorated


def _img(data, key, prefix, company_id):
    """يحفظ صورة base64 إن أُرسلت — اسم الملف يحمل _co<id>_ لتقييد الوصول"""
    b64 = data.get(key)
    if not b64:
        return None
    return save_image(b64, prefix, f"co{company_id}", data.get(key.replace("_base64", "_mime")))


# ════════════════════════════════════════
# هوية الشركة
# ════════════════════════════════════════
@company_bp.route("/api/company/profile", methods=["GET"])
@require_company
def get_profile(account, company_id):
    with get_db() as conn:
        co = conn.execute("SELECT * FROM companies WHERE id=?", (company_id,)).fetchone()
    return jsonify({"company": dict(co)})


@company_bp.route("/api/company/profile", methods=["POST", "PUT"])
@require_company
def save_profile(account, company_id):
    data = request.get_json() or {}
    fields = [f for f in PROFILE_FIELDS if f in data]
    vals = [_clean(data.get(f)) for f in fields]
    for key, col, prefix in [("gerant_cni_recto_base64", "gerant_cni_recto_path", "co_cni_recto"),
                             ("gerant_cni_verso_base64", "gerant_cni_verso_path", "co_cni_verso")]:
        fn = _img(data, key, prefix, company_id)
        if fn:
            fields.append(col); vals.append(fn)
    rc = _clean(data.get("registre_commerce"))
    with get_db() as conn:
        cur = conn.execute("SELECT * FROM companies WHERE id=?", (company_id,)).fetchone()
        live = conn.execute("""SELECT 1 FROM company_hire_requests WHERE company_id=? AND is_current=1
                               AND statut='مقبول'""", (company_id,)).fetchone()
        if live:
            locked = [lbl for k, lbl in [("nom_ar", "اسم الشركة"), ("registre_commerce", "السجل التجاري"),
                                         ("num_agrement", "رقم الاعتماد"), ("date_agrement", "تاريخ الاعتماد")]
                      if k in data and (_clean(data.get(k)) or None) != (cur[k] or None)]
            if locked:
                return jsonify({"error": "لا يمكن تعديل " + "، ".join(locked) +
                                " ما دامت للشركة رخص سائق أجير سارية — راجع الإدارة"}), 400
        if rc and conn.execute("SELECT 1 FROM companies WHERE registre_commerce=? AND id<>?",
                               (rc, company_id)).fetchone():
            return jsonify({"error": "رقم السجل التجاري مستعمل من طرف شركة أخرى"}), 400
        if fields:
            sets = ", ".join(f"{f}=?" for f in fields)
            conn.execute(f"UPDATE companies SET {sets}, updated_at=datetime('now','localtime') WHERE id=?",
                         vals + [company_id])
            conn.commit()
        co = conn.execute("SELECT * FROM companies WHERE id=?", (company_id,)).fetchone()
    return jsonify({"success": True, "company": dict(co)})


# ════════════════════════════════════════
# CRUD عام للمركبات والسائقين
# ════════════════════════════════════════
def _list(table, company_id):
    with get_db() as conn:
        if table == "company_vehicles":
            rows = conn.execute("""
                SELECT v.*, d.nom_ar AS drv_nom_ar, d.prenom_ar AS drv_prenom_ar,
                       d.num_permis AS drv_num_permis
                FROM company_vehicles v LEFT JOIN company_drivers d ON d.id = v.driver_id
                WHERE v.company_id=? ORDER BY v.id DESC""", (company_id,)).fetchall()
        else:
            rows = conn.execute(f"SELECT * FROM {table} WHERE company_id=? ORDER BY id DESC",
                                (company_id,)).fetchall()
    return [dict(r) for r in rows]


def _save(table, fields, images, company_id, item_id=None):
    data = request.get_json() or {}
    cols = [f for f in fields if f in data]
    vals = [_clean(data.get(f)) for f in cols]
    for key, col, prefix in images:
        fn = _img(data, key, prefix, company_id)
        if fn:
            cols.append(col); vals.append(fn)
    with get_db() as conn:
        if item_id is None:
            cols_all = ["company_id"] + cols
            conn.execute(f"INSERT INTO {table} ({','.join(cols_all)}) VALUES ({','.join('?'*len(cols_all))})",
                         [company_id] + vals)
            item_id = conn.execute("SELECT last_insert_rowid() AS i").fetchone()["i"]
        else:
            if not conn.execute(f"SELECT 1 FROM {table} WHERE id=? AND company_id=?",
                                (item_id, company_id)).fetchone():
                return jsonify({"error": "غير موجود"}), 404
            if cols:
                sets = ", ".join(f"{c}=?" for c in cols)
                conn.execute(f"UPDATE {table} SET {sets}, updated_at=datetime('now','localtime') WHERE id=?",
                             vals + [item_id])
        conn.commit()
        row = conn.execute(f"SELECT * FROM {table} WHERE id=?", (item_id,)).fetchone()
    return jsonify({"success": True, "item": dict(row)})


VEH_IMAGES = [("image_carte_grise_base64", "image_carte_grise_path", "co_cg")]
DRV_IMAGES = [("image_permis_recto_base64", "image_permis_recto_path", "co_permis_recto"),
              ("image_permis_verso_base64", "image_permis_verso_path", "co_permis_verso")]


@company_bp.route("/api/company/vehicles", methods=["GET"])
@require_company
def list_vehicles(account, company_id):
    return jsonify({"vehicles": _list("company_vehicles", company_id)})


@company_bp.route("/api/company/vehicles", methods=["POST"])
@require_company
def add_vehicle(account, company_id):
    if not _clean((request.get_json() or {}).get("num_immatriculation")):
        return jsonify({"error": "رقم التسجيل مطلوب"}), 400
    return _save("company_vehicles", VEHICLE_FIELDS, VEH_IMAGES, company_id)


@company_bp.route("/api/company/vehicles/<int:vid>", methods=["PUT"])
@require_company
def edit_vehicle(account, company_id, vid):
    with get_db() as conn:
        if not conn.execute("SELECT 1 FROM company_vehicles WHERE id=? AND company_id=?",
                            (vid, company_id)).fetchone():
            return jsonify({"error": "غير موجود"}), 404
        why = _vehicle_lock(conn, vid)
        if why:
            return jsonify({"error": why}), 400
    return _save("company_vehicles", VEHICLE_FIELDS, VEH_IMAGES, company_id, vid)


def _pending_change(conn, vid):
    return conn.execute("""SELECT * FROM company_requests WHERE vehicle_id=?
                           AND request_type='تغيير_مركبة' AND statut='جديد'""", (vid,)).fetchone()


def _vehicle_lock(conn, vid):
    """المركبة المرتبطة بعقد توظيف ساري لا تُعدَّل مباشرة — التغيير عبر طلب رسمي"""
    if _pending_change(conn, vid):
        return "لهذه المركبة طلب تغيير قيد الدراسة لدى الإدارة"
    v = conn.execute("SELECT driver_id FROM company_vehicles WHERE id=?", (vid,)).fetchone()
    if v and v["driver_id"] and _active_contract(conn, v["driver_id"]):
        return "المركبة مرتبطة بعقد توظيف ساري — استعمل «تغيير المركبة» (طلب رسمي للإدارة)"
    return None


@company_bp.route("/api/company/vehicles/<int:vid>", methods=["DELETE"])
@require_company
def delete_vehicle(account, company_id, vid):
    with get_db() as conn:
        v = conn.execute("SELECT driver_id FROM company_vehicles WHERE id=? AND company_id=?",
                         (vid, company_id)).fetchone()
        if v and v["driver_id"] and _active_contract(conn, v["driver_id"]):
            return jsonify({"error": "لسائق هذه المركبة عقد توظيف ساري — افسخ العقد أولاً"}), 400
        if _pending_change(conn, vid):
            return jsonify({"error": "لهذه المركبة طلب تغيير قيد الدراسة لدى الإدارة"}), 400
        if not v:
            return jsonify({"error": "غير موجود"}), 404
        if conn.execute("""SELECT 1 FROM company_hire_requests WHERE vehicle_id=?
                           UNION SELECT 1 FROM company_requests WHERE vehicle_id=?""", (vid, vid)).fetchone():
            return jsonify({"error": "لا يمكن حذف مركبة لها عقود أو طلبات سابقة (أرشيف) — يبقى في السجل"}), 400
        conn.execute("DELETE FROM company_vehicles WHERE id=? AND company_id=?", (vid, company_id))
        conn.commit()
    return jsonify({"success": True})


@company_bp.route("/api/company/drivers", methods=["GET"])
@require_company
def list_drivers(account, company_id):
    return jsonify({"drivers": _list("company_drivers", company_id)})


def _driver_dup_error(conn, company_id, data, exclude_id=None, current_nin=None):
    """منع تسجيل نفس السائق مرتين في الشركة + التحقق من طول رقم التعريف الوطني (18 رقماً)."""
    nin = "".join(ch for ch in str(data.get("nin") or "") if ch.isdigit())
    if data.get("nin") and len(nin) != 18 and nin != "".join(ch for ch in str(current_nin or "") if ch.isdigit()):
        return f"رقم التعريف الوطني يجب أن يتكوّن من 18 رقماً (المُدخل: {len(nin)}) — صحّحه يدوياً"
    permis = (_clean(data.get("num_permis")) or "").upper()
    for col, val, lbl in (("nin", nin, "رقم التعريف الوطني"), ("UPPER(num_permis)", permis, "رقم رخصة السياقة")):
        if not val:
            continue
        row = conn.execute(f"SELECT nom_ar, prenom_ar FROM company_drivers WHERE company_id=? AND {col}=?"
                           + (" AND id<>?" if exclude_id else ""),
                           (company_id, val) + ((exclude_id,) if exclude_id else ())).fetchone()
        if row:
            return f"هذا السائق مسجّل مسبقاً في شركتك ({(row['nom_ar'] or '')} {(row['prenom_ar'] or '')}) — نفس {lbl}"
    return None


@company_bp.route("/api/company/drivers", methods=["POST"])
@require_company
def add_driver(account, company_id):
    d = request.get_json() or {}
    if not _clean(d.get("nom_ar")) and not _clean(d.get("nom_fr")):
        return jsonify({"error": "لقب السائق مطلوب"}), 400
    with get_db() as conn:
        err = _driver_dup_error(conn, company_id, d)
    if err:
        return jsonify({"error": err}), 400
    return _save("company_drivers", DRIVER_FIELDS, DRV_IMAGES, company_id)


@company_bp.route("/api/company/drivers/<int:did>", methods=["PUT"])
@require_company
def edit_driver(account, company_id, did):
    data = request.get_json() or {}
    with get_db() as conn:
        cur = conn.execute("SELECT * FROM company_drivers WHERE id=? AND company_id=?",
                           (did, company_id)).fetchone()
        if not cur:
            return jsonify({"error": "غير موجود"}), 404
        err = _driver_dup_error(conn, company_id, data, did, cur["nin"])
        if err:
            return jsonify({"error": err}), 400
        if _active_contract(conn, did):
            # بيانات الهوية والرخصة تظهر في العقد ورخصة سائق أجير — لا تُغيَّر أثناء العقد
            LOCK = [("nom_ar", "اللقب"), ("prenom_ar", "الاسم"), ("nin", "NIN"),
                    ("date_naissance", "تاريخ الميلاد"), ("num_permis", "رقم الرخصة"),
                    ("date_delivrance", "تاريخ إصدار الرخصة")]
            changed = [lbl for k, lbl in LOCK if k in data and (_clean(data.get(k)) or None) != (cur[k] or None)]
            if changed:
                return jsonify({"error": "لا يمكن تعديل: " + "، ".join(changed) +
                                " أثناء عقد توظيف ساري — افسخ العقد أولاً"}), 400
    return _save("company_drivers", DRIVER_FIELDS, DRV_IMAGES, company_id, did)


@company_bp.route("/api/company/drivers/<int:did>", methods=["DELETE"])
@require_company
def delete_driver(account, company_id, did):
    with get_db() as conn:
        if not conn.execute("SELECT 1 FROM company_drivers WHERE id=? AND company_id=?",
                            (did, company_id)).fetchone():
            return jsonify({"error": "غير موجود"}), 404
        if _active_contract(conn, did):
            return jsonify({"error": "للسائق عقد توظيف ساري — افسخ العقد أولاً"}), 400
        if conn.execute("SELECT 1 FROM company_hire_requests WHERE driver_id=?", (did,)).fetchone():
            return jsonify({"error": "لا يمكن حذف سائق له عقود سابقة (أرشيف) — يبقى في السجل"}), 400
        conn.execute("UPDATE company_vehicles SET driver_id=NULL WHERE driver_id=? AND company_id=?",
                     (did, company_id))
        conn.execute("DELETE FROM company_drivers WHERE id=? AND company_id=?", (did, company_id))
        conn.commit()
    return jsonify({"success": True})


# ════════════════════════════════════════
# ربط مركبة بسائق (سائق واحد لكل مركبة، ومركبة واحدة لكل سائق)
# ════════════════════════════════════════
@company_bp.route("/api/company/vehicles/<int:vid>/driver", methods=["PUT"])
@require_company
def link_driver(account, company_id, vid):
    did = (request.get_json() or {}).get("driver_id") or None
    with get_db() as conn:
        veh = conn.execute("SELECT driver_id FROM company_vehicles WHERE id=? AND company_id=?",
                           (vid, company_id)).fetchone()
        if not veh:
            return jsonify({"error": "المركبة غير موجودة"}), 404
        # مثل تغيير المناوب: لا يُغيَّر سائق مركبة له عقد ساري إلا بعد فسخ العقد
        if veh["driver_id"] and str(veh["driver_id"]) != str(did or "") and _active_contract(conn, veh["driver_id"]):
            return jsonify({"error": "لسائق هذه المركبة عقد توظيف ساري — افسخ العقد أولاً ثم غيّر السائق"}), 400
        if did and _active_contract(conn, did):
            return jsonify({"error": "لهذا السائق عقد توظيف ساري على مركبة أخرى — افسخ العقد أولاً"}), 400
        if did:
            if not conn.execute("SELECT 1 FROM company_drivers WHERE id=? AND company_id=?",
                                (did, company_id)).fetchone():
                return jsonify({"error": "السائق غير موجود"}), 404
            conn.execute("UPDATE company_vehicles SET driver_id=NULL WHERE driver_id=? AND company_id=?",
                         (did, company_id))
        conn.execute("UPDATE company_vehicles SET driver_id=?, updated_at=datetime('now','localtime') WHERE id=?",
                     (did, vid))
        conn.commit()
    return jsonify({"success": True})



# ════════════════════════════════════════
# عقود توظيف السائقين الأجراء — نفس منطق تبويب المناوب:
#   إنشاء العقد (سنة من اليوم) + إرسال الطلب للإدارة
#   → طباعة طلب التوظيف / العقد
#   → قبول الإدارة = رخصة سائق أجير
#   → فسخ العقد + طباعة المحضر (تُلغى الرخصة) ثم يمكن تغيير السائق
# ════════════════════════════════════════
HIRE_SELECT = """
    SELECT h.*, d.nom_ar AS drv_nom_ar, d.prenom_ar AS drv_prenom_ar,
           d.num_permis AS drv_num_permis, d.date_expiration AS drv_permis_expiration,
           c.nom_ar AS co_nom_ar, c.registre_commerce AS co_rc,
           c.num_agrement AS co_num_agrement, c.date_agrement AS co_date_agrement
    FROM company_hire_requests h
    JOIN company_drivers d ON d.id = h.driver_id
    JOIN companies c       ON c.id = h.company_id
"""


@company_bp.route("/api/company/hire-requests", methods=["GET"])
@require_company
def list_hire_requests(account, company_id):
    with get_db() as conn:
        rows = conn.execute(HIRE_SELECT + " WHERE h.company_id=? ORDER BY h.id DESC",
                            (company_id,)).fetchall()
    return jsonify({"requests": [dict(r) for r in rows]})


@company_bp.route("/api/company/hire-requests", methods=["POST"])
@require_company
def create_hire_contract(account, company_id):
    from dateutil.relativedelta import relativedelta
    did = (request.get_json() or {}).get("driver_id")
    with get_db() as conn:
        co = conn.execute("SELECT * FROM companies WHERE id=?", (company_id,)).fetchone()
        missing = [lbl for k, lbl in [("nom_ar", "اسم الشركة"), ("registre_commerce", "رقم السجل التجاري"),
                                      ("num_agrement", "رقم الاعتماد"), ("date_agrement", "تاريخ الاعتماد")]
                   if not co[k]]
        if missing:
            return jsonify({"error": "أكمل هوية الشركة أولاً: " + "، ".join(missing)}), 400
        drv = conn.execute("SELECT * FROM company_drivers WHERE id=? AND company_id=?",
                           (did, company_id)).fetchone()
        if not drv:
            return jsonify({"error": "السائق غير موجود"}), 404
        if not drv["num_permis"]:
            return jsonify({"error": "رقم رخصة السياقة للسائق غير مُدخل"}), 400
        exp = (drv["date_expiration"] or "")[:10]
        if exp and exp < datetime.now().strftime("%Y-%m-%d"):
            return jsonify({"error": f"رخصة سياقة السائق منتهية الصلاحية ({exp}) — جدّدها أولاً"}), 400
        veh = conn.execute("SELECT * FROM company_vehicles WHERE driver_id=? AND company_id=?",
                           (did, company_id)).fetchone()
        if not veh:
            return jsonify({"error": "اربط السائق بمركبة أولاً قبل إنشاء عقد التوظيف"}), 400
        if _pending_change(conn, veh["id"]):
            return jsonify({"error": "للمركبة طلب تغيير قيد الدراسة — انتظر قرار الإدارة قبل إنشاء العقد"}), 400
        cur = _active_contract(conn, did)
        if cur:
            return jsonify({"error": f"يوجد عقد توظيف ساري ({cur['contract_number']}) — افسخه أولاً ثم أنشئ عقداً جديداً"}), 400
    start = datetime.now().date()
    end = start + relativedelta(years=1)
    number = generate_number("EMP", "company_hire_requests", "contract_number")
    with get_db() as conn:
        c = conn.execute("""INSERT INTO company_hire_requests
            (company_id, driver_id, vehicle_id, num_immatriculation, contract_number,
             contract_start, contract_end)
            VALUES (?,?,?,?,?,?,?)""",
            (company_id, did, veh["id"], veh["num_immatriculation"], number,
             start.isoformat(), end.isoformat()))
        conn.commit()
    return jsonify({"success": True, "id": c.lastrowid, "contract_number": number,
                    "contract_start": start.isoformat(), "contract_end": end.isoformat()})


@company_bp.route("/api/company/hire-requests/<int:rid>/terminate", methods=["POST"])
@require_company
def terminate_hire_contract(account, company_id, rid):
    now = datetime.now()
    with get_db() as conn:
        h = conn.execute("""SELECT * FROM company_hire_requests
                            WHERE id=? AND company_id=? AND is_current=1 AND statut IN ('جديد','مقبول')""",
                         (rid, company_id)).fetchone()
        if not h:
            return jsonify({"error": "لا يوجد عقد توظيف ساري"}), 404
        # فسخ العقد يُلغي رخصة سائق أجير (إن صدرت) ويسحب الطلب إن كان قيد الدراسة
        conn.execute("""UPDATE company_hire_requests
                        SET is_current=0, end_reason='فسخ_يدوي', terminated_at=?,
                            statut = CASE WHEN statut='جديد' THEN 'ملغى' ELSE statut END
                        WHERE id=?""", (now.strftime("%Y-%m-%d %H:%M:%S"), rid))
        conn.commit()
    return jsonify({"success": True})


# ── الوثائق المطبوعة ──
def _fmt(d):
    if not d:
        return "&#8212;"
    d = str(d)[:10]
    if len(d) == 10 and d[4] == "-":
        return f"{d[8:10]}/{d[5:7]}/{d[0:4]}"
    return escape(d)


def _load_hire(rid, company_id=None):
    with get_db() as conn:
        q = "SELECT * FROM company_hire_requests WHERE id=?" + (" AND company_id=?" if company_id else "")
        h = conn.execute(q, (rid, company_id) if company_id else (rid,)).fetchone()
        if not h:
            return None
        h = dict(h)
        d  = dict(conn.execute("SELECT * FROM company_drivers WHERE id=?", (h["driver_id"],)).fetchone() or {})
        co = dict(conn.execute("SELECT * FROM companies WHERE id=?", (h["company_id"],)).fetchone() or {})
        vr = conn.execute("SELECT * FROM company_vehicles WHERE id=?", (h["vehicle_id"],)).fetchone()
        veh = dict(vr) if vr else {"num_immatriculation": h["num_immatriculation"]}
    return h, d, co, veh


def _e(x):
    return escape(str(x)) if x not in (None, "", "None") else "&#8212;"


def _employer_rows(co):
    gerant = f"{co.get('gerant_prenom_ar') or ''} {co.get('gerant_nom_ar') or ''}".strip()
    return f"""
    <tr><td class="label">اسم الشركة</td><td>{_e(co.get("nom_ar"))}</td></tr>
    <tr><td class="label">المقر الاجتماعي</td><td>{_e(co.get("adresse"))} &#8212; {_e(co.get("commune"))} &#8212; {_e(co.get("wilaya"))}</td></tr>
    <tr><td class="label">رقم السجل التجاري</td><td>{_e(co.get("registre_commerce"))}</td></tr>
    <tr><td class="label">رقم التعريف الجبائي</td><td>{_e(co.get("num_fiscal"))}</td></tr>
    <tr><td class="label">رقم وتاريخ الاعتماد</td><td>{_e(co.get("num_agrement"))} &#8212; {_fmt(co.get("date_agrement"))}</td></tr>
    <tr><td class="label">المسيّر</td><td>{_e(gerant)}</td></tr>
    <tr><td class="label">الهاتف</td><td>{_e(co.get("telephone"))}</td></tr>"""


def _driver_rows(d):
    return f"""
    <tr><td class="label">اللقب والاسم</td><td>{_e(d.get("prenom_ar"))} {_e(d.get("nom_ar"))}</td></tr>
    <tr><td class="label">NIN</td><td>{_e(d.get("nin"))}</td></tr>
    <tr><td class="label">تاريخ ومكان الميلاد</td><td>{_fmt(d.get("date_naissance"))} &#8212; {_e(d.get("lieu_naissance"))}</td></tr>
    <tr><td class="label">الهاتف</td><td>{_e(d.get("telephone"))}</td></tr>
    <tr><td class="label">رقم رخصة السياقة</td><td>{_e(d.get("num_permis"))}</td></tr>
    <tr><td class="label">فئات الرخصة</td><td>{_e(d.get("categories"))}</td></tr>
    <tr><td class="label">تاريخ الإصدار / الانتهاء</td><td>{_fmt(d.get("date_delivrance"))} &#8212; {_fmt(d.get("date_expiration"))}</td></tr>"""


def _vehicle_rows(veh):
    return f"""
    <tr><td class="label">رقم التسجيل</td><td>{_e(veh.get("num_immatriculation"))}</td></tr>
    <tr><td class="label">الصنف / الطراز</td><td>{_e(veh.get("marque"))} &#8212; {_e(veh.get("type_vehicule"))}</td></tr>
    <tr><td class="label">الرقم التسلسلي</td><td>{_e(veh.get("num_serie"))}</td></tr>
    <tr><td class="label">الطاقة / عدد المقاعد</td><td>{_e(veh.get("energie"))} &#8212; {_e(veh.get("nb_places"))}</td></tr>"""


# تنسيق طباعة مضغوط — كل وثيقة في صفحة A4 واحدة (مثل وثائق المناوب)
COMPACT_PRINT = """<style>
@media print {
  * { -webkit-print-color-adjust: exact !important; print-color-adjust: exact !important; }
  @page { size: A4 portrait; margin: 8mm; }
  html, body { height: auto !important; margin: 0 !important; padding: 0 !important; }
  body { background: #fff !important; display: block !important; }
  /* الإطار يملأ الورقة تماماً: 297mm - 2×8mm هامش = 281mm (مع هامش أمان صغير) */
  .doc { box-sizing: border-box !important; width: 100% !important; max-width: 100% !important;
         height: 279mm !important; min-height: 0 !important; margin: 0 !important;
         padding: 12px 16px !important; border: 2px solid #125950 !important; border-radius: 6px !important;
         display: flex !important; flex-direction: column !important; overflow: hidden !important; }
  .header { padding-bottom: 6px !important; margin-bottom: 8px !important; }
  .header h1 { font-size: 15px !important; }
  .header p  { font-size: 11px !important; margin-top: 2px !important; }
  .badge, .danger { font-size: 12px !important; padding: 5px 8px !important; margin-bottom: 8px !important; }
  /* الجدول يتمدد ليملأ المساحة وتتوزع الأسطر بالتساوي */
  table { flex: 1 1 auto !important; margin-bottom: 0 !important; width: 100% !important; }
  td { padding: 2px 8px !important; font-size: 12px !important; line-height: 1.3 !important; }
  td.sep, td.sep-red { padding: 3px 8px !important; font-size: 12px !important; }
  td.label { font-size: 12px !important; width: 38% !important; }
  .sign { margin-top: 10px !important; padding-top: 0 !important; flex: 0 0 auto !important; }
  .sign-box .line { margin-top: 45px !important; font-size: 11px !important; }
  .footer { margin-top: 6px !important; font-size: 9px !important; flex: 0 0 auto !important; }
  .btn { display: none !important; }
}
</style>"""


def render_hire_doc(kind, rid, company_id=None):
    from routes.print import html_page
    data = _load_hire(rid, company_id)
    if not data:
        return Response("العقد غير موجود", 404)
    h, d, co, veh = data
    num = _e(h.get("contract_number"))
    today = datetime.now().strftime("%Y-%m-%d")
    sign = """<div class="sign">
      <div class="sign-box"><div class="line">توقيع وختم صاحب العمل (الشركة)</div></div>
      <div class="sign-box"><div class="line">توقيع السائق الأجير</div></div></div>"""
    if kind == "job":
        title = "طلب توظيف سائق أجير"
        body = f"""<div class="doc">
  <div class="header"><h1>&#128661; طلب توظيف سائق أجير لسيارة أجرة</h1><p>محرر بتاريخ: {_fmt(h["contract_start"])}</p></div>
  <div class="badge">رقم العقد: {num}</div>
  <table>
    <tr><td class="sep" colspan="2">&#127970; بيانات الشركة (صاحب العمل)</td></tr>{_employer_rows(co)}
    <tr><td class="sep" colspan="2">&#128100; بيانات السائق الأجير</td></tr>{_driver_rows(d)}
    <tr><td class="sep" colspan="2">&#128664; بيانات المركبة</td></tr>{_vehicle_rows(veh)}
    <tr><td class="sep" colspan="2">&#128203; العقد والنشاط</td></tr>
    <tr><td class="label">رقم عقد العمل</td><td>{num}</td></tr>
    <tr><td class="label">مدة العقد</td><td>سنة واحدة</td></tr>
    <tr><td class="label">تاريخ بداية العقد</td><td>{_fmt(h["contract_start"])}</td></tr>
    <tr><td class="label">تاريخ انتهاء العقد</td><td>{_fmt(h["contract_end"])}</td></tr>
    <tr><td class="label">نوع النشاط</td><td>حضرية فردية</td></tr>
  </table>
  <div class="sign"><div class="sign-box"></div>
    <div class="sign-box"><div class="line">توقيع وختم صاحب العمل (الشركة)</div></div></div>
  <div class="footer">{num} &#8212; {today}</div></div>"""
    elif kind == "contract":
        title = "عقد عمل سائق أجير"
        body = f"""<div class="doc">
  <div class="header"><h1>&#128661; عقد عمل سائق أجير</h1><p>محرر بتاريخ: {_fmt(h["contract_start"])}</p></div>
  <div class="badge">رقم العقد: {num}</div>
  <table>
    <tr><td class="sep" colspan="2">الطرف الأول &#8212; صاحب العمل (الشركة)</td></tr>{_employer_rows(co)}
    <tr><td class="sep" colspan="2">الطرف الثاني &#8212; السائق الأجير</td></tr>{_driver_rows(d)}
    <tr><td class="sep" colspan="2">&#128203; بنود العقد</td></tr>
    <tr><td class="label">مدة العقد</td><td>سنة واحدة</td></tr>
    <tr><td class="label">تاريخ بداية العقد</td><td>{_fmt(h["contract_start"])}</td></tr>
    <tr><td class="label">تاريخ انتهاء العقد</td><td>{_fmt(h["contract_end"])}</td></tr>
    <tr><td class="label">نوع النشاط</td><td>حضرية فردية</td></tr>
    <tr><td class="sep" colspan="2">&#128664; المركبة المكلّف بقيادتها</td></tr>{_vehicle_rows(veh)}
  </table>{sign}
  <div class="footer">{num}</div></div>"""
    else:  # termination
        title = "فسخ عقد سائق أجير"
        term = (h.get("terminated_at") or "")[:10]
        body = f"""<div class="doc">
  <div class="header"><h1>&#128661; محضر فسخ عقد عمل سائق أجير</h1><p>محرر بتاريخ: {_fmt(term or today)}</p></div>
  <div class="danger">&#9888;&#65039; وثيقة فسخ عقد العمل &#8212; رقم العقد: {num}</div>
  <table>
    <tr><td class="sep" colspan="2">الطرف الأول &#8212; صاحب العمل (الشركة)</td></tr>{_employer_rows(co)}
    <tr><td class="sep-red" colspan="2">الطرف الثاني &#8212; السائق الأجير (المفسوخ عقده)</td></tr>{_driver_rows(d)}
    <tr><td class="sep-red" colspan="2">&#128203; بيانات الفسخ</td></tr>
    <tr><td class="label">رقم العقد المفسوخ</td><td>{num}</td></tr>
    <tr><td class="label">المركبة</td><td>{_e(h.get("num_immatriculation"))}</td></tr>
    <tr><td class="label">تاريخ إبرام العقد</td><td>{_fmt(h["contract_start"])}</td></tr>
    <tr><td class="label">تاريخ انتهاء العقد</td><td>{_fmt(h["contract_end"])}</td></tr>
    <tr><td class="label">تاريخ الفسخ</td><td>{_fmt(term) if term else "العقد ساري — لم يُفسخ بعد"}</td></tr>
    <tr><td class="label">رخصة سائق أجير الملغاة</td><td>{_e(h.get("permit_number"))}</td></tr>
  </table>{sign}
  <div class="footer">{num} &#8212; فسخ العقد &#8212; {today}</div></div>"""
    return Response(html_page(title, COMPACT_PRINT + body), mimetype="text/html; charset=utf-8")


def render_hire_permit(rid, company_id=None):
    from routes.admin import DEPUTY_PERMIT_DOC_STYLE
    data = _load_hire(rid, company_id)
    if not data:
        return Response("الطلب غير موجود", 404)
    h, d, co, veh = data
    if h["statut"] != "مقبول":
        return Response("<p dir=rtl style='font-family:sans-serif;padding:30px'>⚠️ لم تُحرَّر الرخصة بعد — الطلب لم يُقبل من طرف الإدارة.</p>",
                        mimetype="text/html; charset=utf-8")
    warn = "" if h["is_current"] else '<div class="warn">&#9888;&#65039; رخصة ملغاة &#8212; تم فسخ عقد التوظيف</div>'
    today = datetime.now().strftime("%d/%m/%Y")
    num = _e(h["permit_number"])
    body = f"""<!DOCTYPE html>
<html dir="rtl" lang="ar"><head><meta charset="UTF-8"><title>رخصة سائق أجير</title>{DEPUTY_PERMIT_DOC_STYLE}</head>
<body><div class="doc">
  <div class="gov-header">
    <div class="g1">الجمهورية الجزائرية الديمقراطية الشعبية</div>
    <div class="g2">وزارة الداخلية والجماعات المحلية والنقل</div>
    <div class="g3">مديرية النقل لولاية البيض</div>
  </div>
  {warn}
  <div class="permit-ref"><div><strong>الرقم:</strong> &nbsp;&nbsp;{num}</div></div>
  <div class="main-title">رخصة سائق أجير</div>
  <div class="body-text" style="margin-bottom:12px;">ان مدير النقل لولاية البيض:</div>
  <div class="body-text">
    - بمقتضى المرسوم التنفيذي 230/12 المؤرخ في 03 رجب عام 1433 الموافق ل 24 مايو 2012
    والمتضمن تنظيم النقل بواسطة سيارات الأجرة المعدل والمتمم.
  </div>
  <div class="body-text" style="margin-bottom:18px;">
    - بمقتضى القرار المؤرخ في 11 ذي القعدة عام 1437 الموافق ل 14 غشت سنة 2016 الذي يحدد نماذج
    الوثائق المرتبطة بممارسة نشاط النقل بواسطة سيارة الأجرة.
  </div>
  <div style="margin-bottom:8px; font-size:13px;">- يقرر ما يأتي-</div>
  <div class="body-text" style="margin-bottom:18px; line-height:2.2;">
    <strong>المادة الأولى :</strong>
    عملا بأحكام المرسوم التنفيذي 230/12 المذكور أعلاه تسلم رخصة سائق أجير
    <strong>للسيد :</strong> {_e(d.get("prenom_ar"))} {_e(d.get("nom_ar"))} ،
    <strong>تاريخ و مكان الازدياد :</strong> <span dir="ltr">{_fmt(d.get("date_naissance"))}</span> &nbsp;ب&nbsp; {_e(d.get("lieu_naissance"))} ،
    <strong>رقم رخصة السياقة :</strong> <span dir="ltr">{_e(d.get("num_permis"))}</span> ،
    <strong>تاريخ رخصة السياقة :</strong> <span dir="ltr">{_fmt(d.get("date_delivrance"))}</span> .
  </div>
  <div class="body-text" style="margin-top:10px; margin-bottom:18px;">
    وذلك تبعا للطلب الذي قدمته شركة : <strong>{_e(co.get("nom_ar"))}</strong> ،
    السجل التجاري رقم : <span dir="ltr">{_e(co.get("registre_commerce"))}</span> ،
    بموجب عقد العمل رقم : <span dir="ltr">{_e(h.get("contract_number"))}</span> ،
    الساري من <span dir="ltr">{_fmt(h["contract_start"])}</span> إلى <span dir="ltr">{_fmt(h["contract_end"])}</span> .
  </div>
  <div class="field-list">
    <div><strong>بلدية الالتحاق :</strong> &nbsp;&nbsp;{_e(co.get("commune") or "البيض")}</div>
    <div><strong>نوع النشاط :</strong> &nbsp;&nbsp;حضرية فردية</div>
    <div><strong>رقم تسجيل المركبة :</strong> &nbsp;&nbsp;<span dir="ltr">{_e(h["num_immatriculation"])}</span></div>
    <div><strong>رقم الاعتماد :</strong> &nbsp;&nbsp;<span dir="ltr">{_e(co.get("num_agrement"))}</span>
         &nbsp;&nbsp; <strong>بتاريخ :</strong> <span dir="ltr">{_fmt(co.get("date_agrement"))}</span></div>
  </div>
  <div style="margin-top:16px; font-size:13px;">
    <strong>&#10005; صالحة الى غاية:</strong> &nbsp;&nbsp;<span dir="ltr">{_fmt(h["expiry_date"])}</span>
  </div>
  <div style="margin-top:40px; text-align:center;">
    <div style="font-size:13px;">حرر بالبيض في : <span dir="ltr">{_fmt(h["issue_date"])}</span></div>
    <div style="font-size:13px; font-weight:700; margin-top:6px;">المدير</div>
  </div>
  <div class="footer">{num} &#8212; رخصة سائق أجير &#8212; {today}</div>
</div>
<button class="btn" onclick="window.print()">&#128424;&#65039; طباعة</button>
</body></html>"""
    return Response(body, mimetype="text/html; charset=utf-8")


DOC_KINDS = {"hire-job-request": "job", "hire-contract": "contract", "hire-termination": "termination"}


@company_bp.route("/api/company/print/<kind>/<int:rid>")
@require_company
def company_print_hire(account, company_id, kind, rid):
    # طباعة رخصة سائق أجير من اختصاص الإدارة فقط
    if kind == "hire-permit":
        return Response("طباعة الرخصة من اختصاص الإدارة", status=403)
    if kind not in DOC_KINDS:
        return Response("غير موجود", 404)
    return render_hire_doc(DOC_KINDS[kind], rid, company_id)


# ── جهة الإدارة ──
@company_bp.route("/api/admin/hire-requests", methods=["GET"])
@require_admin
def admin_list_hire_requests(account):
    with get_db() as _c:
        expire_contracts(_c)
    statut = request.args.get("statut", "")
    with get_db() as conn:
        rows = conn.execute(HIRE_SELECT + (" WHERE h.statut=?" if statut else "") + " ORDER BY h.id DESC",
                            (statut,) if statut else ()).fetchall()
        new = conn.execute("SELECT COUNT(*) FROM company_hire_requests WHERE statut='جديد'").fetchone()[0]
    return jsonify({"requests": [dict(r) for r in rows], "new_count": new})


@company_bp.route("/api/admin/hire-requests/<int:rid>", methods=["PUT"])
@require_admin
def admin_process_hire_request(account, rid):
    with get_db() as _c:
        expire_contracts(_c)
    data = request.get_json() or {}
    statut = data.get("statut")
    if statut not in ("مقبول", "مرفوض"):
        return jsonify({"error": "حالة غير صالحة"}), 400
    notes = _clean(data.get("admin_notes"))
    now = datetime.now()
    with get_db() as conn:
        h = conn.execute("SELECT * FROM company_hire_requests WHERE id=?", (rid,)).fetchone()
        if not h:
            return jsonify({"error": "الطلب غير موجود"}), 404
        if h["statut"] != "جديد" or not h["is_current"]:
            return jsonify({"error": "تمت معالجة هذا الطلب أو فُسخ العقد"}), 400
        if statut == "مرفوض":
            if not notes:
                return jsonify({"error": "اكتب سبب الرفض"}), 400
            conn.execute("""UPDATE company_hire_requests SET statut='مرفوض', is_current=0, admin_notes=?,
                            processed_by=?, processed_at=? WHERE id=?""",
                         (notes, account["id"], now.strftime("%Y-%m-%d %H:%M:%S"), rid))
            conn.commit()
            return jsonify({"success": True})
        drv = conn.execute("SELECT date_expiration FROM company_drivers WHERE id=?", (h["driver_id"],)).fetchone()
    lic = (drv["date_expiration"] or "")[:10] if drv else ""
    today = now.strftime("%Y-%m-%d")
    if lic and lic < today:
        return jsonify({"error": f"رخصة سياقة السائق منتهية ({lic}) — لا يمكن تحرير رخصة سائق أجير"}), 400
    # الرخصة لا تتجاوز نهاية العقد ولا نهاية رخصة السياقة
    expiry = min(x for x in (h["contract_end"], lic) if x)
    permit = generate_number("SAJ", "company_hire_requests", "permit_number")
    with get_db() as conn:
        cur = conn.execute("""UPDATE company_hire_requests SET statut='مقبول', admin_notes=?, permit_number=?,
                        issue_date=?, expiry_date=?, processed_by=?, processed_at=?
                        WHERE id=? AND statut='جديد' AND is_current=1""",
                     (notes, permit, today, expiry,
                      account["id"], now.strftime("%Y-%m-%d %H:%M:%S"), rid))
        conn.commit()
        if cur.rowcount != 1:
            return jsonify({"error": "تمت معالجة هذا الطلب أو فُسخ العقد في الأثناء"}), 409
    return jsonify({"success": True, "permit_number": permit})


@company_bp.route("/api/admin/print/company/<kind>/<int:rid>")
@require_admin
def admin_print_hire(account, kind, rid):
    if kind == "hire-permit":
        return render_hire_permit(rid)
    if kind not in DOC_KINDS:
        return Response("غير موجود", 404)
    return render_hire_doc(DOC_KINDS[kind], rid)


# ════════════════════════════════════════
# لوحة المدير — فضاء شركات سيارات الأجرة
# (إحصائيات، قائمة الشركات وتفاصيلها، الأرشيف، تصدير Excel)
# ════════════════════════════════════════
@company_bp.route("/api/admin/company/stats")
@require_admin
def admin_company_stats(account):
    with get_db() as _c:
        expire_contracts(_c)
    today = datetime.now().strftime("%Y-%m-%d")
    q = lambda conn, sql, *a: conn.execute(sql, a).fetchone()[0]
    with get_db() as conn:
        st = {
            "total_companies":   q(conn, "SELECT COUNT(*) FROM companies c JOIN accounts a ON a.id=c.account_id WHERE a.role='company'"),
            "complete_companies": q(conn, """SELECT COUNT(*) FROM companies WHERE nom_ar IS NOT NULL
                                    AND registre_commerce IS NOT NULL AND num_agrement IS NOT NULL
                                    AND date_agrement IS NOT NULL"""),
            "total_vehicles":    q(conn, "SELECT COUNT(*) FROM company_vehicles"),
            "vehicles_no_driver": q(conn, "SELECT COUNT(*) FROM company_vehicles WHERE driver_id IS NULL"),
            "total_drivers":     q(conn, "SELECT COUNT(*) FROM company_drivers"),
            "expired_permis":    q(conn, "SELECT COUNT(*) FROM company_drivers WHERE date_expiration IS NOT NULL AND date_expiration < ?", today),
            "total_requests":    q(conn, "SELECT COUNT(*) FROM company_hire_requests"),
            "new_requests":      q(conn, "SELECT COUNT(*) FROM company_hire_requests WHERE statut='جديد'"),
            "approved_requests": q(conn, "SELECT COUNT(*) FROM company_hire_requests WHERE statut='مقبول'"),
            "rejected_requests": q(conn, "SELECT COUNT(*) FROM company_hire_requests WHERE statut='مرفوض'"),
            "cancelled_requests": q(conn, "SELECT COUNT(*) FROM company_hire_requests WHERE statut='ملغى'"),
            "today_requests":    q(conn, "SELECT COUNT(*) FROM company_hire_requests WHERE substr(created_at,1,10)=?", today),
            "active_contracts":  q(conn, "SELECT COUNT(*) FROM company_hire_requests WHERE is_current=1 AND statut IN ('جديد','مقبول')"),
            "active_permits":    q(conn, "SELECT COUNT(*) FROM company_hire_requests WHERE is_current=1 AND statut='مقبول'"),
            "terminated":        q(conn, "SELECT COUNT(*) FROM company_hire_requests WHERE end_reason IS NOT NULL"),
            "vehicle_requests_new": q(conn, "SELECT COUNT(*) FROM company_requests WHERE statut='جديد'"),
        }
    return jsonify({"stats": st})


COMPANY_LIST_SQL = """
    SELECT c.*, a.username, a.google_email, a.is_active,
      (SELECT COUNT(*) FROM company_vehicles v WHERE v.company_id=c.id) AS nb_vehicles,
      (SELECT COUNT(*) FROM company_drivers d WHERE d.company_id=c.id) AS nb_drivers,
      (SELECT COUNT(*) FROM company_hire_requests h WHERE h.company_id=c.id
          AND h.is_current=1 AND h.statut='مقبول') AS nb_permits,
      (SELECT COUNT(*) FROM company_hire_requests h WHERE h.company_id=c.id AND h.statut='جديد') AS nb_pending
    FROM companies c JOIN accounts a ON a.id = c.account_id
    WHERE a.role='company'
"""


@company_bp.route("/api/admin/companies")
@require_admin
def admin_list_companies(account):
    with get_db() as _c:
        expire_contracts(_c)
    s = (request.args.get("search") or "").strip()
    sql, args = COMPANY_LIST_SQL, []
    if s:
        sql += """ AND (c.nom_ar LIKE ? OR c.nom_fr LIKE ? OR c.registre_commerce LIKE ?
                   OR c.num_agrement LIKE ? OR a.username LIKE ? OR c.telephone LIKE ?)"""
        args = [f"%{s}%"] * 6
    with get_db() as conn:
        rows = conn.execute(sql + " ORDER BY c.id DESC", args).fetchall()
    return jsonify({"companies": [dict(r) for r in rows]})


@company_bp.route("/api/admin/companies/<int:cid>")
@require_admin
def admin_company_detail(account, cid):
    with get_db() as _c:
        expire_contracts(_c)
    with get_db() as conn:
        co = conn.execute(COMPANY_LIST_SQL + " AND c.id=?", (cid,)).fetchone()
        if not co:
            return jsonify({"error": "الشركة غير موجودة"}), 404
        vehicles = _list("company_vehicles", cid)
        drivers = _list("company_drivers", cid)
        hires = [dict(r) for r in conn.execute(HIRE_SELECT + " WHERE h.company_id=? ORDER BY h.id DESC",
                                              (cid,)).fetchall()]
    return jsonify({"company": dict(co), "vehicles": vehicles, "drivers": drivers, "hires": hires})


@company_bp.route("/api/admin/company/archive")
@require_admin
def admin_company_archive(account):
    with get_db() as _c:
        expire_contracts(_c)
    s = (request.args.get("search") or "").strip()
    sql, args = HIRE_SELECT + " WHERE h.is_current=0", []
    if s:
        sql += " AND (c.nom_ar LIKE ? OR d.nom_ar LIKE ? OR d.prenom_ar LIKE ? OR h.contract_number LIKE ? OR h.num_immatriculation LIKE ?)"
        args = [f"%{s}%"] * 5
    with get_db() as conn:
        rows = conn.execute(sql + " ORDER BY h.id DESC", args).fetchall()
    return jsonify({"requests": [dict(r) for r in rows]})


@company_bp.route("/api/admin/company/export")
@require_admin
def admin_company_export(account):
    with get_db() as _c:
        expire_contracts(_c)
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    sheets = [
        ("الشركات", """SELECT c.nom_ar, c.nom_fr, c.registre_commerce, c.rc_date, c.num_fiscal,
                c.num_agrement, c.date_agrement, c.adresse, c.commune, c.wilaya, c.telephone, c.email,
                trim(coalesce(c.gerant_prenom_ar,'')||' '||coalesce(c.gerant_nom_ar,'')), c.gerant_nin,
                (SELECT COUNT(*) FROM company_vehicles v WHERE v.company_id=c.id),
                (SELECT COUNT(*) FROM company_drivers d WHERE d.company_id=c.id)
             FROM companies c JOIN accounts a ON a.id=c.account_id WHERE a.role='company' ORDER BY c.id""",
         ["اسم الشركة", "Raison sociale", "السجل التجاري", "تاريخ السجل", "NIF", "رقم الاعتماد",
          "تاريخ الاعتماد", "العنوان", "البلدية", "الولاية", "الهاتف", "البريد", "المسيّر", "NIN المسيّر",
          "عدد المركبات", "عدد السائقين"]),
        ("المركبات", """SELECT c.nom_ar, v.num_immatriculation, v.marque, v.type_vehicule, v.num_serie,
                v.energie, v.nb_places, v.annee_circulation,
                trim(coalesce(d.prenom_ar,'')||' '||coalesce(d.nom_ar,''))
             FROM company_vehicles v JOIN companies c ON c.id=v.company_id
             LEFT JOIN company_drivers d ON d.id=v.driver_id ORDER BY c.id, v.id""",
         ["الشركة", "رقم التسجيل", "الصنف", "الطراز", "الرقم التسلسلي", "الطاقة", "المقاعد", "سنة الاستعمال", "السائق"]),
        ("السائقون", """SELECT c.nom_ar, d.nom_ar, d.prenom_ar, d.nin, d.date_naissance, d.lieu_naissance,
                d.telephone, d.num_permis, d.categories, d.date_delivrance, d.date_expiration
             FROM company_drivers d JOIN companies c ON c.id=d.company_id ORDER BY c.id, d.id""",
         ["الشركة", "اللقب", "الاسم", "NIN", "تاريخ الميلاد", "مكان الميلاد", "الهاتف", "رقم الرخصة",
          "الأصناف", "تاريخ الإصدار", "تاريخ الانتهاء"]),
        ("العقود والرخص", """SELECT c.nom_ar, trim(coalesce(d.prenom_ar,'')||' '||coalesce(d.nom_ar,'')),
                h.num_immatriculation, h.contract_number, h.contract_start, h.contract_end, h.statut,
                h.permit_number, h.issue_date, h.expiry_date,
                CASE WHEN h.is_current=1 THEN 'ساري' ELSE 'منتهي' END, substr(h.terminated_at,1,10), h.admin_notes
             FROM company_hire_requests h JOIN companies c ON c.id=h.company_id
             JOIN company_drivers d ON d.id=h.driver_id ORDER BY h.id""",
         ["الشركة", "السائق الأجير", "المركبة", "رقم العقد", "بداية العقد", "نهاية العقد", "حالة الطلب",
          "رقم رخصة سائق أجير", "تاريخ التحرير", "صالحة إلى", "الوضعية", "تاريخ الفسخ", "ملاحظات"]),
    ]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="125950")
    with get_db() as conn:
        for title, sql, cols in sheets:
            ws = wb.create_sheet(title)
            ws.sheet_view.rightToLeft = True
            ws.append(cols)
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_fill
                cell.alignment = Alignment(horizontal="center")
            for row in conn.execute(sql).fetchall():
                ws.append(list(row))
            for i, col in enumerate(cols, 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, len(col) + 4)
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=companies.xlsx"})



# ════════════════════════════════════════
# تغيير مركبة الشركة — نفس مبدأ «تغيير_سيارة» عند السائقين:
#   1) الشركة تُدخل المركبة الجديدة (البطاقة الرمادية إلزامية) ← تُؤرشَف القديمة في الطلب
#   2) يُنشأ طلب «تغيير_مركبة» ويُطبع ويُقدَّم للإدارة
#   3) القبول: يُحدَّث عقد التوظيف الساري برقم المركبة الجديدة (تُعاد طباعة رخصة سائق أجير)
#      الرفض: تُسترجع المركبة القديمة
# ════════════════════════════════════════
import json as _json
VEH_REQUIRED = [("num_immatriculation", "رقم التسجيل"), ("marque", "الصنف"),
                ("num_serie", "الرقم التسلسلي في الطراز"), ("nb_places", "عدد المقاعد")]


def _veh_snapshot(row):
    d = dict(row)
    return {k: d.get(k) for k in VEHICLE_FIELDS + ["image_carte_grise_path", "driver_id"]}


@company_bp.route("/api/company/vehicles/<int:vid>/change", methods=["POST"])
@require_company
def change_vehicle(account, company_id, vid):
    data = request.get_json() or {}
    if not data.get("image_carte_grise_base64"):
        return jsonify({"error": "يجب رفع صورة البطاقة الرمادية للمركبة الجديدة"}), 400
    missing = [lbl for k, lbl in VEH_REQUIRED if not _clean(data.get(k))]
    if missing:
        return jsonify({"error": "الحقول التالية ناقصة: " + "، ".join(missing)}), 400
    with get_db() as conn:
        old = conn.execute("SELECT * FROM company_vehicles WHERE id=? AND company_id=?",
                           (vid, company_id)).fetchone()
        if not old:
            return jsonify({"error": "المركبة غير موجودة"}), 404
        if _pending_change(conn, vid):
            return jsonify({"error": "لهذه المركبة طلب تغيير قيد الدراسة لدى الإدارة"}), 400
        if _clean(data.get("num_immatriculation")) == old["num_immatriculation"]:
            return jsonify({"error": "رقم التسجيل هو نفسه رقم المركبة الحالية"}), 400
    fn = save_image(data["image_carte_grise_base64"], "co_cg", f"co{company_id}", data.get("image_carte_grise_mime"))
    new_vals = {f: _clean(data.get(f)) for f in VEHICLE_FIELDS}
    new_vals["image_carte_grise_path"] = fn
    number = generate_number("VCH", "company_requests", "request_number")
    with get_db() as conn:
        old_snap = _veh_snapshot(old)
        sets = ", ".join(f"{k}=?" for k in new_vals)
        conn.execute(f"UPDATE company_vehicles SET {sets}, updated_at=datetime('now','localtime') WHERE id=?",
                     list(new_vals.values()) + [vid])
        c = conn.execute("""INSERT INTO company_requests (company_id, vehicle_id, request_type, request_number,
                            old_data, new_data) VALUES (?,?,'تغيير_مركبة',?,?,?)""",
                         (company_id, vid, number, _json.dumps(old_snap, ensure_ascii=False),
                          _json.dumps(new_vals, ensure_ascii=False)))
        conn.commit()
    return jsonify({"success": True, "id": c.lastrowid, "request_number": number})


REQ_SELECT = """
    SELECT r.*, c.nom_ar AS co_nom_ar, c.registre_commerce AS co_rc,
           v.driver_id, d.nom_ar AS drv_nom_ar, d.prenom_ar AS drv_prenom_ar
    FROM company_requests r
    JOIN companies c ON c.id = r.company_id
    LEFT JOIN company_vehicles v ON v.id = r.vehicle_id
    LEFT JOIN company_drivers d ON d.id = v.driver_id
"""


def _req_rows(rows):
    out = []
    for r in rows:
        r = dict(r)
        for k in ("old_data", "new_data"):
            try:
                r[k] = _json.loads(r[k] or "{}")
            except ValueError:
                r[k] = {}
        out.append(r)
    return out


@company_bp.route("/api/company/vehicle-requests", methods=["GET"])
@require_company
def list_vehicle_requests(account, company_id):
    with get_db() as conn:
        rows = conn.execute(REQ_SELECT + " WHERE r.company_id=? ORDER BY r.id DESC", (company_id,)).fetchall()
    return jsonify({"requests": _req_rows(rows)})


def render_vehicle_change(rid, company_id=None):
    from routes.print import html_page
    with get_db() as conn:
        q = REQ_SELECT + " WHERE r.id=?" + (" AND r.company_id=?" if company_id else "")
        r = conn.execute(q, (rid, company_id) if company_id else (rid,)).fetchone()
        if not r:
            return Response("الطلب غير موجود", 404)
        r = _req_rows([r])[0]
        co = dict(conn.execute("SELECT * FROM companies WHERE id=?", (r["company_id"],)).fetchone())
    o, n = r["old_data"], r["new_data"]
    rows = lambda x: f"""
    <tr><td class="label">رقم التسجيل</td><td>{_e(x.get("num_immatriculation"))}</td></tr>
    <tr><td class="label">الصنف / الطراز</td><td>{_e(x.get("marque"))} &#8212; {_e(x.get("type_vehicule"))}</td></tr>
    <tr><td class="label">الرقم التسلسلي في الطراز</td><td>{_e(x.get("num_serie"))}</td></tr>
    <tr><td class="label">سنة أول استعمال</td><td>{_e(x.get("annee_circulation"))}</td></tr>
    <tr><td class="label">الطاقة / عدد المقاعد</td><td>{_e(x.get("energie"))} &#8212; {_e(x.get("nb_places"))}</td></tr>"""
    drv = f"{r.get('drv_prenom_ar') or ''} {r.get('drv_nom_ar') or ''}".strip()
    status = {"جديد": "قيد الدراسة", "مقبول": "مقبول", "مرفوض": "مرفوض"}.get(r["statut"], r["statut"])
    body = f"""<div class="doc">
  <div class="header"><h1>&#128664; طلب تغيير مركبة &#8212; شركة سيارات الأجرة</h1><p>محرر بتاريخ: {_fmt(r["created_at"])}</p></div>
  <div class="badge">رقم الطلب: {_e(r["request_number"])} &#8212; الحالة: {status}</div>
  <table>
    <tr><td class="sep" colspan="2">&#127970; بيانات الشركة (صاحبة الطلب)</td></tr>{_employer_rows(co)}
    <tr><td class="label">السائق الأجير المكلّف بالمركبة</td><td>{_e(drv)}</td></tr>
    <tr><td class="sep" colspan="2">&#128664; المركبة الحالية (المستبدَلة)</td></tr>{rows(o)}
    <tr><td class="sep" colspan="2">&#128994; المركبة الجديدة</td></tr>{rows(n)}
  </table>
  <div class="sign"><div class="sign-box"></div>
    <div class="sign-box"><div class="line">توقيع وختم صاحب العمل (الشركة)</div></div></div>
  <div class="footer">{_e(r["request_number"])} &#8212; تغيير مركبة</div></div>"""
    return Response(html_page("طلب تغيير مركبة", COMPACT_PRINT + body), mimetype="text/html; charset=utf-8")


@company_bp.route("/api/company/print/vehicle-change/<int:rid>")
@require_company
def company_print_vehicle_change(account, company_id, rid):
    return render_vehicle_change(rid, company_id)


@company_bp.route("/api/admin/print/company/vehicle-change/<int:rid>")
@require_admin
def admin_print_vehicle_change(account, rid):
    return render_vehicle_change(rid)


@company_bp.route("/api/admin/company/vehicle-requests", methods=["GET"])
@require_admin
def admin_vehicle_requests(account):
    statut = request.args.get("statut", "")
    with get_db() as conn:
        rows = conn.execute(REQ_SELECT + (" WHERE r.statut=?" if statut else "") + " ORDER BY r.id DESC",
                            (statut,) if statut else ()).fetchall()
        out = _req_rows(rows)
        # الرخصة الجديدة المحرّرة عند قبول تغيير المركبة — لطباعتها مباشرة من البطاقة
        import re as _re
        for x in out:
            m = _re.search(r"الرخصة الجديدة\s+(\S+)", x.get("admin_notes") or "")
            if x.get("statut") == "مقبول" and m:
                h = conn.execute("SELECT id, is_current FROM company_hire_requests WHERE permit_number=?", (m.group(1),)).fetchone()
                if h:
                    x["permit_hire_id"], x["permit_number"], x["permit_current"] = h["id"], m.group(1), bool(h["is_current"])
    return jsonify({"requests": out})


@company_bp.route("/api/admin/company/vehicle-requests/<int:rid>", methods=["PUT"])
@require_admin
def admin_process_vehicle_request(account, rid):
    data = request.get_json() or {}
    statut = data.get("statut")
    notes = _clean(data.get("admin_notes"))
    if statut not in ("مقبول", "مرفوض"):
        return jsonify({"error": "حالة غير صالحة"}), 400
    if statut == "مرفوض" and not notes:
        return jsonify({"error": "اكتب سبب الرفض"}), 400
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with get_db() as conn:
        r = conn.execute("SELECT * FROM company_requests WHERE id=?", (rid,)).fetchone()
        if not r:
            return jsonify({"error": "الطلب غير موجود"}), 404
        if r["statut"] != "جديد":
            return jsonify({"error": "تمت معالجة هذا الطلب مسبقاً"}), 400
        new = _json.loads(r["new_data"] or "{}")
        old = _json.loads(r["old_data"] or "{}")
        permit = None
        if statut == "مقبول":
            # عقد التوظيف الساري على هذه المركبة يأخذ رقم المركبة الجديدة → تُعاد طباعة الرخصة
            conn.execute("""UPDATE company_hire_requests SET num_immatriculation=?
                            WHERE vehicle_id=? AND is_current=1 AND statut IN ('جديد','مقبول')""",
                         (new.get("num_immatriculation"), r["vehicle_id"]))
            p = conn.execute("""SELECT id, permit_number FROM company_hire_requests
                                WHERE vehicle_id=? AND is_current=1 AND statut='مقبول'""",
                             (r["vehicle_id"],)).fetchone()
            if p:
                # مثل السائقين: قبول تغيير المركبة = منح رخصة جديدة تتضمن المركبة الجديدة
                # (رقم جديد وتاريخ تحرير جديد؛ الرخصة السابقة تُلغى بالاستبدال)
                conn.commit()
                new_num = generate_number("SAJ", "company_hire_requests", "permit_number")
                conn.execute("UPDATE company_hire_requests SET permit_number=?, issue_date=? WHERE id=?",
                             (new_num, now[:10], p["id"]))
                trace = f"الرخصة السابقة {p['permit_number']} ← الرخصة الجديدة {new_num}"
                notes = f"{notes} — {trace}" if notes else trace
                permit = {"id": p["id"], "permit_number": new_num, "old_permit_number": p["permit_number"]}
        else:
            # الرفض: استرجاع المركبة القديمة (وفي العقد الساري أيضاً)
            conn.execute("""UPDATE company_hire_requests SET num_immatriculation=?
                            WHERE vehicle_id=? AND is_current=1 AND statut IN ('جديد','مقبول')""",
                         (old.get("num_immatriculation"), r["vehicle_id"]))
            cols = [k for k in VEHICLE_FIELDS + ["image_carte_grise_path"] if k in old]
            if cols:
                conn.execute(f"UPDATE company_vehicles SET {', '.join(c + '=?' for c in cols)}, "
                             "updated_at=datetime('now','localtime') WHERE id=?",
                             [old[c] for c in cols] + [r["vehicle_id"]])
        conn.execute("""UPDATE company_requests SET statut=?, admin_notes=?, processed_by=?, processed_at=?
                        WHERE id=?""", (statut, notes, account["id"], now, rid))
        conn.commit()
    return jsonify({"success": True, "permit": permit})



# ════════════════════════════════════════
# بطاقة معلومات الشركة + الشهادة التاريخية للشركة
#   الشركة: تطبع وثائقها هي فقط — الإدارة: أي شركة
# ════════════════════════════════════════
def _company_bundle(cid):
    with get_db() as conn:
        expire_contracts(conn)
        co = conn.execute("""SELECT c.*, a.username, a.created_at AS acc_created FROM companies c
                             JOIN accounts a ON a.id=c.account_id WHERE c.id=?""", (cid,)).fetchone()
        if not co:
            return None
        vehicles = [dict(r) for r in conn.execute("""
            SELECT v.*, d.nom_ar AS drv_nom_ar, d.prenom_ar AS drv_prenom_ar
            FROM company_vehicles v LEFT JOIN company_drivers d ON d.id=v.driver_id
            WHERE v.company_id=? ORDER BY v.id""", (cid,)).fetchall()]
        drivers = [dict(r) for r in conn.execute(
            "SELECT * FROM company_drivers WHERE company_id=? ORDER BY id", (cid,)).fetchall()]
        hires = [dict(r) for r in conn.execute("""
            SELECT h.*, d.nom_ar AS drv_nom_ar, d.prenom_ar AS drv_prenom_ar, d.nin AS drv_nin
            FROM company_hire_requests h LEFT JOIN company_drivers d ON d.id=h.driver_id
            WHERE h.company_id=? ORDER BY h.id""", (cid,)).fetchall()]
        vreqs = [dict(r) for r in conn.execute(
            "SELECT * FROM company_requests WHERE company_id=? ORDER BY id", (cid,)).fetchall()]
    for r in vreqs:
        for k in ("old_data", "new_data"):
            try:
                r[k] = _json.loads(r[k] or "{}")
            except ValueError:
                r[k] = {}
    return dict(co), vehicles, drivers, hires, vreqs


def _join(*xs, sep=" &#8212; "):
    parts = [_e(x) for x in xs if x not in (None, "", "None")]
    return sep.join(parts) if parts else "&#8212;"


def _co_identity_rows(co):
    gerant = f"{co.get('gerant_prenom_ar') or ''} {co.get('gerant_nom_ar') or ''}".strip()
    return f"""
    <tr><td class="sep" colspan="2">&#127970; هوية الشركة</td></tr>
    <tr><td class="label">اسم الشركة</td><td>{_e(co.get("nom_ar"))}</td></tr>
    <tr><td class="label">Raison sociale</td><td dir="ltr" style="text-align:right">{_e(co.get("nom_fr"))}</td></tr>
    <tr><td class="label">المقر الاجتماعي</td><td>{_join(co.get("adresse"), co.get("commune"), co.get("wilaya"))}</td></tr>
    <tr><td class="label">رقم وتاريخ السجل التجاري</td><td>{_e(co.get("registre_commerce"))}{(" &#8212; " + _fmt(co.get("rc_date"))) if co.get("rc_date") else ""}</td></tr>
    <tr><td class="label">رقم التعريف الجبائي</td><td>{_e(co.get("num_fiscal"))}</td></tr>
    <tr><td class="label">رقم وتاريخ الاعتماد</td><td>{_e(co.get("num_agrement"))}{(" &#8212; " + _fmt(co.get("date_agrement"))) if co.get("date_agrement") else ""}</td></tr>
    <tr><td class="label">الهاتف</td><td>{_e(co.get("telephone"))}{(" &#8212; " + _e(co.get("telephone2"))) if co.get("telephone2") else ""}</td></tr>
    <tr><td class="label">البريد الإلكتروني</td><td>{_e(co.get("email"))}</td></tr>
    <tr><td class="sep" colspan="2">&#128100; المسيّر</td></tr>
    <tr><td class="label">اللقب والاسم</td><td>{_e(gerant)}</td></tr>
    <tr><td class="label">تاريخ ومكان الميلاد</td><td>{_join(_fmt(co.get("gerant_date_naissance")) if co.get("gerant_date_naissance") else None, co.get("gerant_lieu_naissance"))}</td></tr>
    <tr><td class="label">NIN</td><td>{_e(co.get("gerant_nin"))}</td></tr>
    <tr><td class="label">العنوان</td><td>{_e(co.get("gerant_adresse"))}</td></tr>"""


_H = 'style="background:#f4f8f7;font-weight:700"'
_CUR = ' style="background:#e4f5ec"'
_NONE = '<tr><td colspan="{n}" style="text-align:center;color:#6b7280">لا توجد سجلات</td></tr>'


def render_company_card(cid):
    from routes.print import html_page
    b = _company_bundle(cid)
    if not b:
        return Response("الشركة غير موجودة", 404)
    co, vehicles, drivers, hires, _ = b
    today = datetime.now().strftime("%Y-%m-%d")
    active = {h["driver_id"]: h for h in hires if h["is_current"] and h["statut"] in ("جديد", "مقبول")}

    veh_rows = ""
    for v in vehicles:
        h = active.get(v.get("driver_id"))
        drv = f"{_e(v.get('drv_prenom_ar'))} {_e(v.get('drv_nom_ar'))}" if v.get("driver_id") else "بدون سائق"
        veh_rows += (f"<tr><td dir='ltr'>{_e(v.get('num_immatriculation'))}</td>"
                     f"<td>{_e(v.get('marque'))} {_e(v.get('type_vehicule'))}</td><td>{_e(v.get('num_serie'))}</td>"
                     f"<td>{_e(v.get('nb_places'))}</td><td>{drv}</td>"
                     f"<td>{_e(h.get('permit_number')) if h and h['statut']=='مقبول' else ('قيد الدراسة' if h else '&#8212;')}</td></tr>")

    drv_rows = ""
    for d in drivers:
        h = active.get(d["id"])
        exp = (d.get("date_expiration") or "")[:10]
        exp_cell = _fmt(exp) + (" <b style='color:#b91c1c'>(منتهية)</b>" if exp and exp < today else "")
        if h:
            ct = f"{_e(h.get('contract_number'))} &#8212; إلى {_fmt(h.get('contract_end'))}"
            st = f"رخصة {_e(h.get('permit_number'))} صالحة إلى {_fmt(h.get('expiry_date'))}" if h["statut"] == "مقبول" else "طلب قيد الدراسة"
        else:
            ct, st = "لا عقد ساري", "&#8212;"
        drv_rows += (f"<tr{_CUR if h else ''}><td>{_e(d.get('prenom_ar'))} {_e(d.get('nom_ar'))}</td>"
                     f"<td dir='ltr'>{_e(d.get('nin'))}</td><td dir='ltr'>{_e(d.get('num_permis'))}</td>"
                     f"<td>{exp_cell}</td><td>{ct}</td><td>{st}</td></tr>")

    n_permits = sum(1 for h in active.values() if h["statut"] == "مقبول")
    body = f"""<div class="doc">
    <div class="header"><h1>&#127970; بطاقة معلومات شركة سيارات الأجرة</h1><p>الوضعية الحالية &#8212; صادرة بتاريخ: {_fmt(today)}</p></div>
    <table>{_co_identity_rows(co)}
      <tr><td class="sep" colspan="2">&#128202; ملخص</td></tr>
      <tr><td class="label">عدد المركبات</td><td>{len(vehicles)}</td></tr>
      <tr><td class="label">عدد السائقين الأجراء</td><td>{len(drivers)}</td></tr>
    </table>
    <h3>&#128664; المركبات</h3>
    <table><tr {_H}><td>رقم التسجيل</td><td>الصنف/الطراز</td><td>الرقم التسلسلي</td><td>المقاعد</td><td>السائق</td><td>رخصة سائق أجير</td></tr>
    {veh_rows or _NONE.format(n=6)}</table>
    <h3>&#128100; السائقون الأجراء</h3>
    <table><tr {_H}><td>السائق</td><td>NIN</td><td>رخصة السياقة</td><td>صالحة إلى</td><td>عقد العمل</td><td>رخصة سائق أجير</td></tr>
    {drv_rows or _NONE.format(n=6)}</table>
    <div class="sign">
      <div class="sign-box"><div class="line">توقيع وختم المسيّر</div></div>
      <div class="sign-box"><div class="line">تاريخ الإصدار: {_fmt(today)}</div></div>
    </div>
    </div>"""
    return Response(html_page("بطاقة معلومات الشركة", body), mimetype="text/html; charset=utf-8")


def render_company_history(cid):
    from routes.print import html_page
    b = _company_bundle(cid)
    if not b:
        return Response("الشركة غير موجودة", 404)
    co, vehicles, drivers, hires, vreqs = b
    today = datetime.now().strftime("%Y-%m-%d")
    END = {"انتهاء_المدة": "انتهاء المدة", "فسخ_يدوي": "فسخ"}

    # ── عقود التوظيف ورخص سائق أجير ──
    hire_rows = ""
    for h in hires:
        if h["is_current"] and h["statut"] in ("جديد", "مقبول"):
            st = "ساري &#9989;" if h["statut"] == "مقبول" else "قيد الدراسة"
        elif h["statut"] == "مرفوض":
            st = f"مرفوض &#8212; {_e(h.get('admin_notes'))}"
        elif h["statut"] == "ملغى":
            st = "ملغى (فُسخ قبل القرار)"
        else:
            st = f"منتهٍ &#8212; {END.get(h.get('end_reason'), _e(h.get('end_reason')))} {_fmt((h.get('terminated_at') or '')[:10])}"
        permit = (f"{_e(h.get('permit_number'))}<br><small>{_fmt(h.get('issue_date'))} &#8594; {_fmt(h.get('expiry_date'))}</small>"
                  if h.get("permit_number") else "&#8212;")
        hire_rows += (f"<tr{_CUR if (h['is_current'] and h['statut']=='مقبول') else ''}>"
                      f"<td>{_e(h.get('drv_prenom_ar'))} {_e(h.get('drv_nom_ar'))}</td>"
                      f"<td dir='ltr'>{_e(h.get('num_immatriculation'))}</td><td>{_e(h.get('contract_number'))}</td>"
                      f"<td>{_fmt(h.get('contract_start'))}</td><td>{_fmt(h.get('contract_end'))}</td>"
                      f"<td>{permit}</td><td>{st}</td></tr>")

    # ── تغييرات المركبات ──
    vch_rows = ""
    STATUT = {"جديد": "قيد الدراسة", "مقبول": "مقبول &#9989;", "مرفوض": "مرفوض (استُرجعت القديمة)"}
    for r in vreqs:
        o, n = r["old_data"], r["new_data"]
        vch_rows += (f"<tr><td>{_e(r.get('request_number'))}</td><td>{_fmt((r.get('created_at') or '')[:10])}</td>"
                     f"<td dir='ltr'>{_e(o.get('num_immatriculation'))}</td><td dir='ltr'>{_e(n.get('num_immatriculation'))}</td>"
                     f"<td>{STATUT.get(r['statut'], _e(r['statut']))}"
                     f"{(' &#8212; ' + _fmt((r.get('processed_at') or '')[:10])) if r.get('processed_at') else ''}</td>"
                     f"<td>{_e(r.get('admin_notes'))}</td></tr>")

    # ── سجل السائقين الأجراء (أول وآخر عقد) ──
    by_drv = {}
    for h in hires:
        by_drv.setdefault(h["driver_id"], []).append(h)
    drv_rows = ""
    for d in drivers:
        hs = by_drv.get(d["id"], [])
        cur = any(h["is_current"] and h["statut"] in ("جديد", "مقبول") for h in hs)
        first = _fmt(min((h["contract_start"] for h in hs if h.get("contract_start")), default=None))
        drv_rows += (f"<tr{_CUR if cur else ''}><td>{_e(d.get('prenom_ar'))} {_e(d.get('nom_ar'))}</td>"
                     f"<td dir='ltr'>{_e(d.get('nin'))}</td><td>{_fmt((d.get('created_at') or '')[:10])}</td>"
                     f"<td>{len(hs)}</td><td>{first}</td><td>{'يعمل حالياً &#9989;' if cur else 'لا عقد ساري'}</td></tr>")

    nb = lambda f: sum(1 for h in hires if f(h))
    body = f"""<div class="doc">
    <div class="header"><h1>&#128220; الشهادة التاريخية للشركة</h1><p>شركة سيارات الأجرة &#8212; صادرة بتاريخ: {_fmt(today)}</p></div>
    <table>{_co_identity_rows(co)}
      <tr><td class="sep" colspan="2">&#128202; ملخص المسار</td></tr>
      <tr><td class="label">تاريخ التسجيل في النظام</td><td>{_fmt((co.get("created_at") or co.get("acc_created") or "")[:10])}</td></tr>
      <tr><td class="label">عقود التوظيف</td><td>{len(hires)} (سارية: {nb(lambda h: h["is_current"] and h["statut"] in ("جديد","مقبول"))} &#8212; منتهية/مفسوخة: {nb(lambda h: not h["is_current"] and h["statut"]=="مقبول")} &#8212; مرفوضة: {nb(lambda h: h["statut"]=="مرفوض")} &#8212; ملغاة: {nb(lambda h: h["statut"]=="ملغى")})</td></tr>
      <tr><td class="label">رخص سائق أجير محرَّرة</td><td>{nb(lambda h: h.get("permit_number"))}</td></tr>
      <tr><td class="label">طلبات تغيير المركبات</td><td>{len(vreqs)}</td></tr>
      <tr><td class="label">المركبات / السائقون حالياً</td><td>{len(vehicles)} / {len(drivers)}</td></tr>
    </table>
    <h3>&#128203; عقود التوظيف ورخص سائق أجير</h3>
    <table><tr {_H}><td>السائق</td><td>المركبة</td><td>رقم العقد</td><td>من</td><td>إلى</td><td>رخصة سائق أجير</td><td>الحالة</td></tr>
    {hire_rows or _NONE.format(n=7)}</table>
    <h3>&#128664; تاريخ تغيير المركبات</h3>
    <table><tr {_H}><td>رقم الطلب</td><td>التاريخ</td><td>المركبة السابقة</td><td>المركبة الجديدة</td><td>القرار</td><td>ملاحظات / تتبّع الرخصة</td></tr>
    {vch_rows or _NONE.format(n=6)}</table>
    <h3>&#128100; سجل السائقين الأجراء</h3>
    <table><tr {_H}><td>السائق</td><td>NIN</td><td>تاريخ التسجيل</td><td>عدد العقود</td><td>أول عقد</td><td>الوضعية</td></tr>
    {drv_rows or _NONE.format(n=6)}</table>
    <div class="sign">
      <div class="sign-box"><div class="line">المدير &#8212; ختم وتوقيع</div></div>
      <div class="sign-box"><div class="line">تاريخ الإصدار: {_fmt(today)}</div></div>
    </div>
    </div>"""
    return Response(html_page("الشهادة التاريخية للشركة", body), mimetype="text/html; charset=utf-8")


@company_bp.route("/api/company/print-company/<doc>")
@require_company
def company_print_own(account, company_id, doc):
    if doc == "card":
        return render_company_card(company_id)
    # الشهادة التاريخية وثيقة إدارية — تُطبع من لوحة الإدارة (مثل السائقين)
    return Response("غير موجود", 404)


@company_bp.route("/api/admin/print-company/<doc>/<int:cid>")
@require_admin
def admin_print_company(account, doc, cid):
    if doc == "card":
        return render_company_card(cid)
    if doc == "history":
        return render_company_history(cid)
    return Response("غير موجود", 404)
